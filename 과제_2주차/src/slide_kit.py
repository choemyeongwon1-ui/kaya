"""슬라이드 제작 공용 헬퍼 (build_slide.py, build_assignment.py에서 사용)."""
from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FONT = "Malgun Gothic"
NAVY = RGBColor(0x1F, 0x3A, 0x68)
BLUE = RGBColor(0x1F, 0x4E, 0x9C)
INK = RGBColor(0x26, 0x2B, 0x33)
MUTED = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF5, 0xF7, 0xFA)
HEADER = RGBColor(0xF2, 0xF4, 0xF7)
LINE = RGBColor(0xE4, 0xE8, 0xEF)


def style_run(run, size=12, color=INK, bold=False, italic=False, spacing=None):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "ko-KR")
    etree.SubElement(rPr, qn("a:ea")).set("typeface", FONT)
    if spacing:
        rPr.set("spc", str(spacing))


def text(slide, x, y, w, h, runs, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, spacing=None, line_spacing=None):
    """runs: str 또는 [(텍스트, {옵션})] 리스트. '\\n'으로 문단 구분."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [(runs, {})]
    p = tf.paragraphs[0]
    p.alignment = align
    p._p.get_or_add_pPr().set("eaLnBrk", "0")  # 한글 단어가 줄 끝에서 잘리지 않게
    if line_spacing:
        p.line_spacing = line_spacing
    for chunk, opt in runs:
        for i, part in enumerate(chunk.split("\n")):
            if i:
                p = tf.add_paragraph()
                p.alignment = align
                p._p.get_or_add_pPr().set("eaLnBrk", "0")
                if line_spacing:
                    p.line_spacing = line_spacing
            if not part:
                continue
            style_run(p.add_run(), size=opt.get("size", size), color=opt.get("color", color),
                      bold=opt.get("bold", bold), italic=opt.get("italic", italic),
                      spacing=spacing)
            p.runs[-1].text = part
            if opt.get("link"):  # 클릭하면 열리는 주소
                p.runs[-1].hyperlink.address = opt["link"]
    return tb


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None, line_w=1.25):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s


def label_shape(shape, label, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    p._p.get_or_add_pPr().set("eaLnBrk", "0")
    style_run(p.add_run(), size=size, color=color, bold=bold)
    p.runs[-1].text = label
    return shape


def slide_header(slide, section, title, page, course="스마트도시계획(캡스톤디자인)"):
    rect(slide, 0, 0, 13.333, 1.5, HEADER)
    rect(slide, 0.45, 0.3, 0.1, 0.32, BLUE)
    text(slide, 0.7, 0.27, 7.5, 0.38, section, size=14, bold=True, color=BLUE, spacing=120,
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 0.6, 0.75, 10.5, 0.55, title, size=28, bold=True, color=NAVY)
    text(slide, 8.4, 0.27, 3.95, 0.38, f"{course} |", size=11, color=INK, align=PP_ALIGN.RIGHT,
         anchor=MSO_ANCHOR.MIDDLE)
    label_shape(rect(slide, 12.45, 0.22, 0.5, 0.5, BLUE), str(page), size=14)


def slide_footer(slide, source, y=7.0):
    text(slide, 0.6, y, 12.1, 0.3, source, size=9.5, italic=True, color=MUTED)
